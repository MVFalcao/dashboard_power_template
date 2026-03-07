from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook


def build_workbook(output_path: Path, total_rows: int = 50) -> None:
    workbook = Workbook()
    ws_candidatos = workbook.active
    ws_candidatos.title = "Candidatos"

    ws_aprovados = workbook.create_sheet("Aprovados")
    ws_analise = workbook.create_sheet("Em análise")
    ws_negado = workbook.create_sheet("Negado")
    ws_database = workbook.create_sheet("Database")
    ws_legenda = workbook.create_sheet("legenda")

    candidatos_headers = [
        "Nome do inscrito(a)",
        "Data de inscrição",
        "e-mail",
        "Contato",
        "Região",
        "Turma",
        "Idade",
        "Renda Familiar",
        "Equipe NAU",
        "Agendamento",
        "Status",
        "Motivo",
        "Comentarios",
    ]

    status_sheet_headers = [
        "Nome do inscrito(a)",
        "Data de inscrição",
        "e-mail",
        "Contato",
        "Região",
        "Turma",
        "Idade",
        "Renda Familiar",
        "Equipe NAU",
        "Agendamento",
        "Motivo",
        "Comentarios",
    ]

    database_headers = [
        "Data de inscrição",
        "Nome do inscrito(a)",
        "Idade",
        "e-mail",
        "Celular/WhatsApp",
        "Região",
        "Turma",
        "Renda Familiar",
    ]

    for sheet in [ws_candidatos]:
        sheet.append(candidatos_headers)

    for sheet in [ws_aprovados, ws_analise, ws_negado]:
        sheet.append(status_sheet_headers)

    ws_database.append(database_headers)
    ws_legenda.append(["Região", "Status", "Equipe ", "Turma", "Idade", "Motivo", "Renda Familiar"])

    regions = ["SP", "BA", "PE", "RJ", "AM", "MG"]
    turmas = ["Online T", "Online M", "Presencial"]
    idades = ["Menos de 16 anos", "Entre 16 a 24 anos", "Mais de 24 anos"]
    rendas = ["Até 1 SM", "1-2 SM", "2-4 SM", "Acima de 4 SM"]
    equipes = ["pessoa1", "pessoa2", "pessoa3", "pessoa4"]
    status_cycle = ["Aprovado", "Negado", "Em análise"]

    motivos_map = {
        "Aprovado": ["Perfil aderente", "Disponibilidade confirmada"],
        "Negado": ["Conflito de agenda", "Sem requisitos"],
        "Em análise": ["Documentação pendente", "Aguardando entrevista"],
    }

    base_date = date(2026, 1, 1)

    for idx in range(1, total_rows + 1):
        inscrição = (base_date + timedelta(days=idx - 1)).isoformat()
        nome = f"Candidato {idx:02d}"
        email = f"candidato{idx:02d}@teste.com"
        telefone = f"+55 11 9{idx:08d}"[-14:]
        regiao = regions[(idx - 1) % len(regions)]
        turma = turmas[(idx - 1) % len(turmas)]
        idade = idades[(idx - 1) % len(idades)]
        renda = rendas[(idx - 1) % len(rendas)]
        equipe = equipes[(idx - 1) % len(equipes)]
        status = status_cycle[(idx - 1) % len(status_cycle)]
        motivo = motivos_map[status][(idx - 1) % len(motivos_map[status])]
        comentario = f"Registro de teste {idx:02d}"
        agendamento = (base_date + timedelta(days=idx + 7)).isoformat()

        ws_database.append([inscrição, nome, idade, email, telefone, regiao, turma, renda])

        candidatos_row = [
            nome,
            inscrição,
            email,
            telefone,
            regiao,
            turma,
            idade,
            renda,
            equipe,
            agendamento,
            status,
            motivo,
            comentario,
        ]
        ws_candidatos.append(candidatos_row)

        status_row = [
            nome,
            inscrição,
            email,
            telefone,
            regiao,
            turma,
            idade,
            renda,
            equipe,
            agendamento,
            motivo,
            comentario,
        ]

        if status == "Aprovado":
            ws_aprovados.append(status_row)
        elif status == "Negado":
            ws_negado.append(status_row)
        else:
            ws_analise.append(status_row)

    ws_legenda_rows = [
        ["SP", "Aprovado", "pessoa1", "Online T", "Menos de 16 anos", "Perfil aderente", "Até 1 SM"],
        ["BA", "Negado", "pessoa2", "Online M", "Entre 16 a 24 anos", "Conflito de agenda", "1-2 SM"],
        ["PE", "Em análise", "pessoa3", "Presencial", "Mais de 24 anos", "Documentação pendente", "2-4 SM"],
    ]
    for row in ws_legenda_rows:
        ws_legenda.append(row)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


if __name__ == "__main__":
    target = Path("template") / "Dashboard_50.xlsx"
    build_workbook(target, total_rows=50)
    print(f"arquivo_gerado={target}")
