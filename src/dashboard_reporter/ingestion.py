from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import AppConfig
from .models import IngestionResult, RawRow
from .utils import file_sha256, is_empty, stable_hash_from_mapping


def _sanitize_header_value(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _normalize_cell_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def _extract_headers(worksheet, header_row: int) -> dict[int, str]:
    headers: dict[int, str] = {}
    used_names: dict[str, int] = {}

    for column_idx in range(1, worksheet.max_column + 1):
        raw = worksheet.cell(header_row, column_idx).value
        name = _sanitize_header_value(raw)
        if not name:
            continue

        if name in used_names:
            used_names[name] += 1
            name = f"{name}_{used_names[name]}"
        else:
            used_names[name] = 1
        headers[column_idx] = name

    return headers


def ingest_workbook(workbook_path: Path, config: AppConfig) -> IngestionResult:
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook não encontrado: {workbook_path}")

    wb_values = load_workbook(workbook_path, data_only=True)
    wb_formulas = load_workbook(workbook_path, data_only=False)

    raw_rows: list[RawRow] = []
    sheet_stats: dict[str, dict[str, int]] = {}

    for sheet_name in config.source.input_sheets:
        if sheet_name not in wb_values.sheetnames:
            sheet_stats[sheet_name] = {
                "missing_sheet": 1,
                "headers": 0,
                "rows_kept": 0,
                "rows_skipped_empty": 0,
            }
            continue

        ws_values = wb_values[sheet_name]
        ws_formulas = wb_formulas[sheet_name]
        headers = _extract_headers(ws_values, config.source.header_row)

        if not headers:
            sheet_stats[sheet_name] = {
                "missing_sheet": 0,
                "headers": 0,
                "rows_kept": 0,
                "rows_skipped_empty": max(ws_values.max_row - config.source.header_row, 0),
            }
            continue

        kept = 0
        skipped = 0

        for row_idx in range(config.source.header_row + 1, ws_values.max_row + 1):
            values: dict[str, Any] = {}
            formulas: dict[str, str] = {}
            non_empty = False

            for col_idx, header in headers.items():
                value_cell = ws_values.cell(row_idx, col_idx).value
                formula_cell = ws_formulas.cell(row_idx, col_idx).value

                normalized_value = _normalize_cell_value(value_cell)
                values[header] = normalized_value

                if isinstance(formula_cell, str) and formula_cell.startswith("="):
                    formulas[header] = formula_cell

                if not is_empty(normalized_value):
                    non_empty = True

            if not non_empty:
                skipped += 1
                continue

            raw_rows.append(
                RawRow(
                    sheet_name=sheet_name,
                    row_number=row_idx,
                    values=values,
                    formulas=formulas,
                    source_file=str(workbook_path),
                    row_hash=stable_hash_from_mapping({"sheet": sheet_name, "row": row_idx, "values": values}),
                )
            )
            kept += 1

        sheet_stats[sheet_name] = {
            "missing_sheet": 0,
            "headers": len(headers),
            "rows_kept": kept,
            "rows_skipped_empty": skipped,
        }

    return IngestionResult(
        source_file=str(workbook_path),
        file_hash=file_sha256(workbook_path),
        raw_rows=raw_rows,
        sheet_stats=sheet_stats,
    )
